import os
import sys
import json
import io
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Force UTF-8 encoding on Windows
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

from src.mining.models import CompanyRecord
from src.mining.international_miner import InternationalManufacturingMiner
from src.mining.deduplicator import CompanyDeduplicator
from src.mining.validator import CompanyDataValidator
from src.mining.exporter import (
    DeliverableExporter,
    HEADER_FONT, HEADER_FILL, SUBHEADER_FILL, ACCENT_FILL, TOTAL_FONT, BORDER_THIN, SCHEMA_FIELDS
)

def update_excel_with_international_data():
    print("======================================================================")
    print("  UPDATING EXISTING EXCEL & MASTER DELIVERABLES")
    print("  Manufacturing Customer Engagement: US, India & Asia (Japan, Malaysia, Thailand)")
    print("======================================================================\n")

    # 1. Load existing records
    existing_json = "output/master_companies.json"
    if not os.path.exists(existing_json):
        print(f"[!] {existing_json} not found. Running from checkpoint or base.")
        existing_json = "output/extraction_checkpoint.json"

    with open(existing_json, "r", encoding="utf-8") as f:
        existing_raw = json.load(f)

    print(f"[*] Loaded {len(existing_raw)} existing company records.")
    existing_records = [CompanyRecord(**r) for r in existing_raw]

    # 2. Mine International Manufacturing Customer Engagement Records
    miner = InternationalManufacturingMiner()
    new_records = miner.run_mining()

    # 3. Combine and Deduplicate
    all_raw = existing_records + new_records
    print(f"\n[*] Combining datasets: {len(all_raw)} total raw records before deduplication.")
    unique_records, dedup_stats = CompanyDeduplicator.deduplicate_and_merge(all_raw)
    print(f"[✓] Deduplicated: {len(unique_records)} unique companies, {dedup_stats['duplicates_merged_count']} merged.")

    # 4. Run Quality Validation
    # Support target countries
    validator = CompanyDataValidator(target_states=[])
    val_summary = validator.run_validation_suite(unique_records)
    print(f"[✓] Validation: {val_summary['status_counts']}")

    # 5. Export Updated Deliverables
    exporter = DeliverableExporter("output")

    # A. Update master_companies.csv
    csv_path = exporter.export_csv(unique_records, "master_companies.csv")
    shutil.copy(csv_path, "master_companies.csv")
    print(f"[✓] Updated CSV: output/master_companies.csv and master_companies.csv")

    # B. Update master_companies.json
    json_path = exporter.export_json(unique_records, "master_companies.json")
    print(f"[✓] Updated JSON: output/master_companies.json")

    # C. Update Model JSONs
    model_filenames = {
        "MODEL_01": "model_01_industrial_robotics.json",
        "MODEL_02": "model_02_industrial_automation.json",
        "MODEL_03": "model_03_robotic_components.json",
        "MODEL_04": "model_04_motion_control.json",
    }
    for mid, fname in model_filenames.items():
        exporter.export_model_json(unique_records, mid, fname)
    print(f"[✓] Updated Model JSON files.")

    # D. Build Enhanced Excel Workbook with Country_Summary and State_Summary
    xlsx_path = "output/master_companies.xlsx"
    wb = openpyxl.Workbook()

    # Sheet 1: Master_Data
    ws_master = wb.active
    ws_master.title = "Master_Data"
    ws_master.append(SCHEMA_FIELDS)
    for col_idx in range(1, len(SCHEMA_FIELDS) + 1):
        cell = ws_master.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, rec in enumerate(unique_records, start=2):
        d = rec.to_export_dict()
        row_vals = [d[field] for field in SCHEMA_FIELDS]
        ws_master.append(row_vals)
        for c_idx in range(1, len(row_vals) + 1):
            ws_master.cell(row=row_idx, column=c_idx).border = BORDER_THIN

    # Sheet 2: Country_Summary (US, India, Japan, Malaysia, Thailand, etc.)
    ws_country = wb.create_sheet(title="Country_Summary")
    c_headers = ["Country / Region", "Unique Companies", "Share of Total", "Engagement Scope"]
    ws_country.append(c_headers)
    for c in range(1, len(c_headers) + 1):
        cell = ws_country.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    country_counts = {}
    for r in unique_records:
        ctry = r.country or "Unknown"
        country_counts[ctry] = country_counts.get(ctry, 0) + 1

    total_unique = len(unique_records)
    row_num = 2
    for ctry, cnt in sorted(country_counts.items(), key=lambda x: x[1], reverse=True):
        pct = f"{(cnt / total_unique * 100):.1f}%"
        scope = "US Target States & Automation Leaders" if ctry == "United States" else "Asia Manufacturing Customer Engagement"
        ws_country.append([ctry, cnt, pct, scope])
        for c in range(1, 5):
            ws_country.cell(row=row_num, column=c).border = BORDER_THIN
        row_num += 1

    # Country Total row
    ws_country.append(["Total Enterprises", total_unique, "100.0%", "Global Manufacturing Engagement"])
    for c in range(1, 5):
        cell = ws_country.cell(row=row_num, column=c)
        cell.font = TOTAL_FONT
        cell.fill = ACCENT_FILL
        cell.border = BORDER_THIN

    # Sheet 3: State_Summary (State/Province/Prefecture)
    ws_state = wb.create_sheet(title="State_Summary")
    state_headers = ["State / Province / Prefecture", "Country", "Unique Companies", "Share"]
    ws_state.append(state_headers)
    for c in range(1, len(state_headers) + 1):
        cell = ws_state.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    state_tuples = {}
    for r in unique_records:
        key = (r.state or "Unknown", r.country or "Unknown")
        state_tuples[key] = state_tuples.get(key, 0) + 1

    row_num = 2
    for (st, ctry), cnt in sorted(state_tuples.items(), key=lambda x: x[1], reverse=True):
        pct = f"{(cnt / total_unique * 100):.1f}%"
        ws_state.append([st, ctry, cnt, pct])
        for c in range(1, 5):
            ws_state.cell(row=row_num, column=c).border = BORDER_THIN
        row_num += 1

    # Sheet 4: Filter_Summary
    ws_filter = wb.create_sheet(title="Filter_Summary")
    f_headers = ["Model / Filter Scope", "Model Name", "Category Domain", "Records Collected", "Filter Applied"]
    ws_filter.append(f_headers)
    for c in range(1, len(f_headers) + 1):
        cell = ws_filter.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    model_meta = [
        ("MODEL_01", "Industrial Robotics", "Robotics, Cobots & Turnkey Cells", "A3 searchTC=2 + Intl Robotics"),
        ("MODEL_02", "Industrial Automation", "Factory Automation, PLCs & Customer Engagement", "A3 searchTC='' + Intl Automation"),
        ("MODEL_03", "Robotic Components", "Sensors, Vision Systems & Grippers", "A3 searchTC=3 + Intl Components"),
        ("MODEL_04", "Motion Control", "Servo Systems, AMRs & Precision Motion", "A3 searchTC=1 + Intl Motion"),
    ]
    row_num = 2
    for mid, mname, cdomain, fmech in model_meta:
        mcount = sum(1 for r in unique_records if r.model_id == mid or mid in getattr(r, "associated_models", []))
        ws_filter.append([mid, mname, cdomain, mcount, fmech])
        for c in range(1, 6):
            ws_filter.cell(row=row_num, column=c).border = BORDER_THIN
        row_num += 1

    # Sheet 5: Category_Summary
    ws_cat = wb.create_sheet(title="Category_Summary")
    cat_headers = ["Industry Category", "Company Count", "Share"]
    ws_cat.append(cat_headers)
    for c in range(1, len(cat_headers) + 1):
        cell = ws_cat.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    cat_counts = {}
    for r in unique_records:
        cat = r.company_category or "Unassigned"
        cat_counts[cat] = cat_counts.get(cat, 0) + 1

    row_num = 2
    for cat, cnt in sorted(cat_counts.items(), key=lambda x: x[1], reverse=True):
        pct = f"{(cnt / total_unique * 100):.1f}%"
        ws_cat.append([cat, cnt, pct])
        for c in range(1, 4):
            ws_cat.cell(row=row_num, column=c).border = BORDER_THIN
        row_num += 1

    # Sheet 6: Validation
    ws_val = wb.create_sheet(title="Validation")
    val_headers = ["Metric / Audit Check", "Count", "Percentage / Details"]
    ws_val.append(val_headers)
    for c in range(1, len(val_headers) + 1):
        cell = ws_val.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    val_stats = val_summary.get("status_counts", {})
    missing = val_summary.get("missing_fields", {})
    raw_cnt = dedup_stats.get("raw_records_count", total_unique)
    merged_cnt = dedup_stats.get("duplicates_merged_count", 0)

    val_rows = [
        ("Total Raw Records Extracted", raw_cnt, "Before deduplication"),
        ("Total Unique Companies", total_unique, "After normalization & deduplication"),
        ("Duplicate Records Merged", merged_cnt, f"{(merged_cnt / max(raw_cnt, 1) * 100):.1f}% duplicate rate"),
        ("Status: VERIFIED", val_stats.get("VERIFIED", 0), f"{(val_stats.get('VERIFIED', 0) / max(total_unique, 1) * 100):.1f}%"),
        ("Status: PARTIAL", val_stats.get("PARTIAL", 0), f"{(val_stats.get('PARTIAL', 0) / max(total_unique, 1) * 100):.1f}%"),
        ("Status: REVIEW", val_stats.get("REVIEW", 0), f"{(val_stats.get('REVIEW', 0) / max(total_unique, 1) * 100):.1f}%"),
        ("Missing Official Website", missing.get("official_website", 0), f"{(missing.get('official_website', 0) / max(total_unique, 1) * 100):.1f}%"),
        ("Missing Public Contact Phone", missing.get("contact_information", 0), f"{(missing.get('contact_information', 0) / max(total_unique, 1) * 100):.1f}%"),
    ]
    row_num = 2
    for metric, cnt, dtl in val_rows:
        ws_val.append([metric, cnt, dtl])
        for c in range(1, 4):
            ws_val.cell(row=row_num, column=c).border = BORDER_THIN
        row_num += 1

    # Sheet 7: Sources
    ws_src = wb.create_sheet(title="Sources")
    src_headers = ["Source Organization / Portal", "Directory / Website URL", "Geographic Scope", "Compliance & Method"]
    ws_src.append(src_headers)
    for c in range(1, len(src_headers) + 1):
        cell = ws_src.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    src_rows = [
        ("A3 - Association for Advancing Automation", "https://www.automate.org/companies/", "United States & Global", "Form filling with startRow offset pagination"),
        ("Official Corporate Portals (India)", "https://wipropari.com/, https://addverb.com/, etc.", "India", "Live HTTP corporate metadata and customer engagement extraction"),
        ("Official Corporate Portals (Japan)", "https://fanuc.co.jp, https://yaskawa-global.com, etc.", "Japan", "Live HTTP corporate metadata and customer engagement extraction"),
        ("Official Corporate Portals (Malaysia)", "https://pentamaster.com.my, https://genetec.net, etc.", "Malaysia", "Live HTTP corporate metadata and customer engagement extraction"),
        ("Official Corporate Portals (Thailand)", "https://deltathailand.com, https://sumipol.com, etc.", "Thailand", "Live HTTP corporate metadata and customer engagement extraction"),
    ]
    row_num = 2
    for srow in src_rows:
        ws_src.append(srow)
        for c in range(1, len(srow) + 1):
            ws_src.cell(row=row_num, column=c).border = BORDER_THIN
        row_num += 1

    # Auto-adjust column widths
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                max_len = max(max_len, min(len(val), 60))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(xlsx_path)
    shutil.copy(xlsx_path, "master_companies.xlsx")
    print(f"[✓] Successfully updated Excel: output/master_companies.xlsx and master_companies.xlsx (7 worksheets)")

    # E. Update validation_report.xlsx
    exporter.export_validation_report_excel(val_summary, dedup_stats, "validation_report.xlsx")
    print(f"[✓] Updated QA Report: output/validation_report.xlsx")

    print("\n======================================================================")
    print(f"  FINAL DATASET SUMMARY")
    print("======================================================================")
    print(f"  Total Unique Companies: {total_unique}")
    print("  Country Distribution:")
    for ctry, cnt in sorted(country_counts.items(), key=lambda x: x[1], reverse=True):
        print(f"    - {ctry}: {cnt} companies")
    print("======================================================================\n")

if __name__ == "__main__":
    update_excel_with_international_data()
